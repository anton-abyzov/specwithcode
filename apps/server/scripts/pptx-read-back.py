"""Reads a PPTX with python-pptx and prints what it holds as JSON.

The design export tests run this to prove the file opens in an independent
reader: slide size, and per slide its text, speaker notes and shape kinds.
Usage: python -I pptx-read-back.py deck.pptx
"""

import json
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def shape_kind(shape):
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if shape.shape_type == MSO_SHAPE_TYPE.FREEFORM:
        return "freeform"
    if shape.has_text_frame and shape.text_frame.text != "":
        return "text"
    return "shape"


def main(path):
    deck = Presentation(path)
    slides = []
    for slide in deck.slides:
        texts = [
            shape.text_frame.text
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text != ""
        ]
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        shapes = [{"name": shape.name, "kind": shape_kind(shape)} for shape in slide.shapes]
        slides.append({"texts": texts, "notes": notes, "shapes": shapes})
    print(
        json.dumps(
            {
                "slideWidth": int(deck.slide_width),
                "slideHeight": int(deck.slide_height),
                "slides": slides,
            }
        )
    )


if __name__ == "__main__":
    main(sys.argv[1])
